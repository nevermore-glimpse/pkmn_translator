# -*- coding: utf-8 -*-
r"""
译前 / 译后替换（借鉴 LinguaGacha Replacement 的**预设**组织方式）。

规则字段
--------
    {"src": "要被替换的内容", "dst": "替换成",
     "regex": false,            # true 时 src 按正则处理，dst 可用 \1 反向引用
     "info": "备注（只给人看）"}

作用位置
--------
· 译前替换：送模型之前 —— 把游戏里的特殊写法归一化；
· 译后替换：译文落盘之前 —— 统一译文排版（如整句首尾半角引号转直角引号）。

预设目录（一个文件一个预设，文件名前缀决定执行顺序）
----------------------------------------------------
    replacements/pre/01_常用.json
    replacements/post/01_中文排版.json
    replacements/post/02_别的.json …

★ 序号前缀只是给人排序看的，实际按**文件名排序**拼接；
  想控制先后就把前缀写对（01_、02_…）。
★ 启用哪些：`PRE/POST_REPLACEMENT_PRESETS` 留空 = 全部启用；
  填 `"01_常用,02_别的"` = 只启用这两个（按填写顺序）。
★ 旧版单文件（`PRE/POST_REPLACEMENT_FILE` 指向具体 json）仍然有效，
  配了就只用那一个文件；没配时若预设目录为空，也回落到
  `replacements/<kind>_default.json`。
★ 预设可导出成 **JSON + Excel 两份**（列一致），也能从任一格式导入。
"""
import json
import os
import re

import config
from logger import get_logger

log = get_logger("replacement")

RULE_KEYS = ("src", "dst", "regex", "info")

# pre / post 各自的预设子目录
PRESET_SUBDIR = {"pre": "pre", "post": "post"}


# ================================================================
# 路径
# ================================================================
def _presets_root(kind):
    """预设根目录：replacements/<pre|post>。"""
    return config.resource_path(
        os.path.join("replacements", PRESET_SUBDIR[kind]))


def _legacy_path(kind):
    """旧版单文件：replacements/<kind>_default.json。"""
    return config.resource_path(
        os.path.join("replacements", f"{kind}_default.json"))


def _user_file(kind):
    return (getattr(config, "PRE_REPLACEMENT_FILE", "") if kind == "pre"
            else getattr(config, "POST_REPLACEMENT_FILE", ""))


def _enabled_names(kind):
    """
    返回「要启用的预设名」列表；留空表示全部启用（返回 None）。
    """
    raw = (getattr(config, "PRE_REPLACEMENT_PRESETS", "") if kind == "pre"
           else getattr(config, "POST_REPLACEMENT_PRESETS", ""))
    names = [n.strip() for n in str(raw or "").replace("；", ",").split(",")
             if n.strip()]
    return names or None


# ================================================================
# 规则归一
# ================================================================
def _norm_rule(item):
    if isinstance(item, str):
        return {"src": item, "dst": "", "regex": False, "info": ""}
    if not isinstance(item, dict):
        return None
    src = str(item.get("src") or "")
    if not src:
        return None
    return {
        "src": src,
        "dst": str(item.get("dst") or ""),
        "regex": bool(item.get("regex")),
        "info": str(item.get("info") or ""),
    }


def _norm_rules(items):
    out = []
    for it in items or []:
        r = _norm_rule(it)
        if r:
            out.append(r)
    return out


# ================================================================
# 读写
# ================================================================
_CACHE = {}       # 路径 → (mtime, size, rules)


def read_rules_file(path):
    """
    读一个规则文件（.json）。文件不存在 / 读不出来 → []。
    带 (mtime, size) 缓存，运行中改文件会生效。
    """
    if not path or not os.path.isfile(path):
        return []
    try:
        st = os.stat(path)
        key = (path, st.st_mtime, st.st_size)
    except OSError:
        return []
    cached = _CACHE.get(path)
    if cached and cached[0] == key:
        return cached[1]

    rules = []
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict) and "rules" in data:
            data = data.get("rules")
        if isinstance(data, list):
            rules = _norm_rules(data)
        elif isinstance(data, dict):
            # 兼容 {原文: 译文} 简写
            rules = _norm_rules([{"src": k, "dst": v}
                                 for k, v in data.items()])
        else:
            log.warning("替换规则文件格式不对（要 JSON 数组）：%s", path)
    except Exception as e:
        log.warning("替换规则文件读取失败（按无规则处理）：%s：%s", path, e)
        rules = []

    _CACHE[path] = (key, rules)
    return rules


def list_presets(kind):
    """列出某类预设：[{"name", "file", "count"}, ...]（按文件名排序）。"""
    d = _presets_root(kind)
    out = []
    if not os.path.isdir(d):
        return out
    for fn in sorted(os.listdir(d)):
        if not fn.lower().endswith(".json"):
            continue
        p = os.path.join(d, fn)
        out.append({
            "name": os.path.splitext(fn)[0],
            "file": p,
            "count": len(read_rules_file(p)),
        })
    return out


def load_rules(kind):
    """
    读一类替换规则（合并当前启用的预设）。

    优先级：
      ① PRE/POST_REPLACEMENT_FILE 指向的文件（配了就只用它）
      ② 预设目录下启用的预设（按文件名排序；PRESETS 名单可筛选）
      ③ 预设目录为空时，回落到旧版 replacements/<kind>_default.json
    """
    user = _user_file(kind)
    if user and os.path.isfile(user):
        rules = read_rules_file(user)
        log.info("替换规则[%s]：%d 条 ← %s（指定文件）",
                 kind, len(rules), user)
        return rules

    presets = list_presets(kind)
    if presets:
        want = _enabled_names(kind)
        if want is not None:
            by_name = {p["name"]: p for p in presets}
            chosen = [by_name[n] for n in want if n in by_name]
            missing = [n for n in want if n not in by_name]
            if missing:
                log.warning("替换预设不存在，已忽略：%s", "、".join(missing))
        else:
            chosen = presets
        rules, used = [], []
        for p in chosen:
            rs = read_rules_file(p["file"])
            if rs:
                used.append(f"{p['name']}({len(rs)})")
            rules += rs
        log.info("替换规则[%s]：%d 条 ← %s", kind, len(rules),
                 "、".join(used) or "（空）")
        return rules

    legacy = _legacy_path(kind)
    if os.path.isfile(legacy):
        rules = read_rules_file(legacy)
        log.info("替换规则[%s]：%d 条 ← %s（旧版单文件）",
                 kind, len(rules), legacy)
        return rules
    return []


def save_preset(kind, name, rules):
    """写一个预设（覆盖同名）。返回写入路径。"""
    d = _presets_root(kind)
    os.makedirs(d, exist_ok=True)
    safe = re.sub(r'[\\/:*?"<>|]', "_", str(name or "").strip()) or "preset"
    path = os.path.join(d, safe + ".json")
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8", newline="") as f:
        json.dump(_norm_rules(rules), f, ensure_ascii=False, indent=2)
    os.replace(tmp, path)
    _CACHE.pop(path, None)
    log.info("替换预设已保存：%s（%d 条）", path, len(_norm_rules(rules)))
    return path


# ================================================================
# Excel 导入 / 导出（缺 openpyxl 时只走 JSON）
# ================================================================
_XLSX_HEADER = ["src", "dst", "regex", "info"]


def export_preset_files(kind, name, rules, base_path=None):
    """
    同时导出 JSON 和 Excel 两份（列一致）。返回 (json_path, xlsx_path|None)。
    Excel 需要 openpyxl，缺了只导出 JSON。
    """
    rules = _norm_rules(rules)
    if base_path:
        base = base_path
    else:
        d = _presets_root(kind)
        os.makedirs(d, exist_ok=True)
        safe = re.sub(r'[\\/:*?"<>|]', "_", str(name or "").strip()) or "preset"
        base = os.path.join(d, safe)

    json_path = base + ".json"
    with open(json_path, "w", encoding="utf-8", newline="") as f:
        json.dump(rules, f, ensure_ascii=False, indent=2)
    _CACHE.pop(json_path, None)

    xlsx_path = None
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "rules"
        ws.append(_XLSX_HEADER)
        for r in rules:
            ws.append([r["src"], r["dst"], bool(r["regex"]), r["info"]])
        xlsx_path = base + ".xlsx"
        wb.save(xlsx_path)
    except ImportError:
        log.warning("没装 openpyxl，本次只导出 JSON")
    except Exception as e:
        log.warning("Excel 导出失败（JSON 已导出）：%s", e)
    return json_path, xlsx_path


def import_rules_file(path):
    """
    从 .json / .xlsx 导入规则（列：src / dst / regex / info）。返回规则列表。
    """
    if not path or not os.path.isfile(path):
        return []
    low = path.lower()
    if low.endswith(".json"):
        return read_rules_file(path)
    if low.endswith(".xlsx"):
        try:
            import openpyxl
        except ImportError:
            log.warning("没装 openpyxl，无法导入 Excel：%s", path)
            return []
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            out = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                cells = ["" if c is None else str(c) for c in row]
                if i == 0 and cells and cells[0].strip().lower() == "src":
                    continue                     # 表头
                if not cells or not cells[0].strip():
                    continue
                out.append({
                    "src": cells[0],
                    "dst": cells[1] if len(cells) > 1 else "",
                    "regex": str(cells[2]).strip().lower() in
                             ("1", "true", "yes", "y"),
                    "info": cells[3] if len(cells) > 3 else "",
                })
            wb.close()
            return _norm_rules(out)
        except Exception as e:
            log.warning("Excel 导入失败：%s：%s", path, e)
            return []
    log.warning("不支持的规则文件格式（只要 .json / .xlsx）：%s", path)
    return []


# ================================================================
# 应用
# ================================================================
def apply_rules(text, rules):
    """按顺序应用规则，返回 (新文本, 命中条数)。"""
    if not text or not rules:
        return text, 0
    out = text
    hits = 0
    for r in rules:
        try:
            if r["regex"]:
                # dst 里的 \1 等反向引用要留给 re.sub 解释，不做转义
                new = re.sub(r["src"], r["dst"], out)
            else:
                new = out.replace(r["src"], r["dst"])
        except re.error as e:
            log.warning("替换规则正则无效，已跳过：%r（%s）", r["src"], e)
            continue
        if new != out:
            hits += 1
            out = new
    return out, hits


def apply_pre(text):
    """译前替换（关闭或没配规则时原样返回）。"""
    if not getattr(config, "PRE_REPLACEMENT_ENABLE", False):
        return text, []
    return apply_rules(text, load_rules("pre"))


def apply_post(text):
    """译后替换。"""
    if not getattr(config, "POST_REPLACEMENT_ENABLE", True):
        return text, []
    return apply_rules(text, load_rules("post"))
