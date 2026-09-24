# -*- coding: utf-8 -*-
"""恢复被重复插入损坏的使用手册：先精确删除插入块，再以正确顺序重新插入 v1.2.0 内容。可重复运行。"""
import shutil, os
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, "dist/使用手册.docx")
CORRUPT = os.path.join(ROOT, "dist/使用手册.corrupted.bak.docx")

# 从损坏备份恢复，保证可重复运行
if os.path.exists(CORRUPT):
    shutil.copy(CORRUPT, SRC)
    print("restored from corrupted backup")
else:
    shutil.copy(SRC, CORRUPT)
    print("saved corrupted copy ->", CORRUPT)

doc = Document(SRC)
VALID = {s.name for s in doc.styles}

def find(text):
    for p in doc.paragraphs:
        if p.text.strip() == text:
            return p
    return None

def remove_paragraph_elements(elements):
    for el in elements:
        el.getparent().remove(el)
    return len(elements)

# ---------- 1) 删除已插入的块 ----------
# TOC 重复条目（精确文本）
for t in ["升级到 v1.2.0（变更说明）", "v1.2.0 新增功能详解"]:
    els = [p._element for p in list(doc.paragraphs) if p.text.strip() == t]
    remove_paragraph_elements(els)
print("removed TOC duplicate entries")

# 第 0 章（从首个 "0. 升级到" 到 "1. 环境要求" 之前，自动覆盖重复副本）
ps = doc.paragraphs
s = next((i for i, p in enumerate(ps) if p.text.strip() == "0. 升级到 v1.2.0（变更说明）"), None)
e = next((i for i, p in enumerate(ps) if p.text.strip() == "1. 环境要求"), None)
if s is not None and e is not None:
    print("removed chapter0:", remove_paragraph_elements([ps[i]._element for i in range(s, e)]))

# 第 14 章（从首个以 "14." 开头的段落到 "文档结束" 之前）
ps = doc.paragraphs
s14 = next((i for i, p in enumerate(ps) if p.text.startswith("14.")), None)
e14 = next((i for i, p in enumerate(ps) if p.text.strip() == "文档结束"), None)
if s14 is not None and e14 is not None:
    print("removed chapter14:", remove_paragraph_elements([ps[i]._element for i in range(s14, e14)]))

# ---------- 2) 还原版本号到 v1.1.0 ----------
def set_runs(p, old, new):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
for p in doc.paragraphs:
    set_runs(p, "v1.2.0", "v1.1.0")
print("reverted version to v1.1.0")

# ---------- 3) 参考推进式插入（正向顺序 -> 文档正向） ----------
def insert_block(anchor, items, before=True):
    ref = None
    for style, text in items:
        new = OxmlElement("w:p")
        if ref is None:
            (anchor._p.addprevious(new) if before else anchor._p.addnext(new))
        else:
            ref._p.addnext(new)
        p = Paragraph(new, anchor._parent)
        p.add_run(text)
        if style in VALID:
            try:
                p.style = style
            except Exception:
                pass
        ref = p

toc = find("常见问题")
insert_block(toc, [
    ("ds-markdown-paragraph", "升级到 v1.2.0（变更说明）"),
    ("ds-markdown-paragraph", "v1.2.0 新增功能详解"),
], before=False)

env = find("1. 环境要求")
changelog = [
    ("heading 2", "0. 升级到 v1.2.0（变更说明）"),
    ("ds-markdown-paragraph", "本版相较于上一版本，在界面、翻译流程与术语管理上做了大量升级，并重点加固了占位符保留。以下为变更摘要，详细用法见第 14 章。"),
    ("ds-markdown-paragraph", "【新增功能】"),
    ("ds-markdown-paragraph", "• 亚克力毛玻璃 GUI：圆形头像、宝可梦主题配色、8 个菜单页、右侧文件侧边栏、状态条小精灵球，支持本地 Ollama 与云端 API 双模式切换。"),
    ("ds-markdown-paragraph", "• 文件范围选择：菜单 1/2/3/5 执行前可选「单文件 / 整个文件夹」，自动跳过生成物，选择结果在菜单间自动同步。"),
    ("ds-markdown-paragraph", "• 中文润色重翻：新增菜单 5，对缓存中的中文译文做一轮润色并覆盖回写。"),
    ("ds-markdown-paragraph", "• 云端 API 模式：新增 OpenAI 兼容 /chat/completions 云端模式，切换时同步调整参数并弹窗提示。"),
    ("ds-markdown-paragraph", "• 检查更新：从 GitHub Release 拉取最新 tag，版本不一致时弹窗提醒（含下载链接）。"),
    ("ds-markdown-paragraph", "• 前缀 / 检查报告列表滚动条：溢出时显示滚动条，长文件不溢出。"),
    ("ds-markdown-paragraph", "【修复与改进】"),
    ("ds-markdown-paragraph", "• 占位符保留加固（重点）：术语表加载排除 @0@ 等占位符 token，杜绝其被当术语注入 prompt 反向教模型翻译占位符；清理了模型误提取的 @0@…@4@ 伪术语；auto_terms 与 @@T: 提取增加占位符型 key 拒绝。"),
    ("ds-markdown-paragraph", "• 译后占位符校验失败时，针对性重试并明确告知模型丢失了哪些 @N@；兜底补回改用源文相对位置精确回插，并修正此前兜底结果被二次判失败、整条被丢弃的潜在缺陷。"),
    ("ds-markdown-paragraph", "• 术语提取 / 匹配性能优化（单词哈希 + 多词小正则）；中文字体统一为微软雅黑，字号适度放大不溢出；译文缺失 / 未保留的控制码按原文原样补回。"),
    ("ds-markdown-paragraph", "【升级说明】直接覆盖运行即可；首次使用请在「设置」中配置模型，云端模式需填写 API 地址与密钥。术语表 term_dict.py 为本地私有数据，不随仓库分发。"),
]
insert_block(env, changelog, before=True)

end = find("文档结束")
chapter14 = [
    ("heading 2", "14. v1.2.0 新增功能详解"),
    ("heading 3", "14.1 毛玻璃 GUI 界面"),
    ("ds-markdown-paragraph", "启动后进入亚克力磨砂质感 GUI：左上圆形头像（悬停显示作者信息），标题栏显示版本 v1.2.0；左侧 8 个菜单，右侧文件侧边栏，底部状态条带小精灵球。"),
    ("ds-markdown-paragraph", "GUI 菜单顺序：1 翻译 / 2 重翻检查报告 / 3 术语更新后重翻 / 4 前缀字典 / 5 中文润色 / 6 Excel 转术语表 / 7 设置 / 8 日志·环境·检查更新。"),
    ("ds-markdown-paragraph", "菜单 8 为日志页，翻译等任务启动后会自动跳转过去查看进度。"),
    ("heading 3", "14.2 文件范围选择"),
    ("ds-markdown-paragraph", "菜单 1/2/3/5 执行前，右侧侧边栏可选择「单个文件」或「整个文件夹」。"),
    ("ds-markdown-paragraph", "选文件夹时列出全部 .txt 并逐个翻译；自动跳过生成物（*_translated.txt、*_translated_report.txt）。"),
    ("ds-markdown-paragraph", "在菜单 1 选择的结果会自动同步到菜单 2~5，无需重复选择。"),
    ("heading 3", "14.3 中文润色（菜单 5，新增）"),
    ("ds-markdown-paragraph", "对缓存中的中文译文再做一轮润色并覆盖回写，不改变原意与占位符。"),
    ("ds-markdown-paragraph", "短于阈值的句子自动跳过；润色时严格校验占位符，不全则保留原译文。"),
    ("heading 3", "14.4 云端 API 模式（新增）"),
    ("ds-markdown-paragraph", "除本地 Ollama 外，新增 OpenAI 兼容 /chat/completions 云端模式；在「设置」或模式切换处选择「云端 API」并填写 API 地址与密钥即可。"),
    ("ds-markdown-paragraph", "切换时自动调整术语上限 / 超时 / 上下文 / 最大生成 / 每批条数 / 单批字符，并弹窗提示前后值。"),
    ("heading 3", "14.5 检查更新（菜单 8，新增）"),
    ("ds-markdown-paragraph", "从 GitHub Release 拉取最新 tag，与本地版本不一致时弹窗提醒并提供下载链接。"),
    ("heading 3", "14.6 术语删除与撤回（菜单 3，新增）"),
    ("ds-markdown-paragraph", "术语列表支持双击编辑译文；可选中（可多选）或按当前筛选批量删除术语。"),
    ("ds-markdown-paragraph", "删除在「应用术语」之前可一键「撤回」，撤回栈支持逐步恢复，便于安心整理术语表。"),
    ("ds-markdown-paragraph", "应用后：改动 / 删除写入术语字典，并删除含相关术语句子的缓存以重翻。"),
]
insert_block(end, chapter14, before=True)

# ---------- 4) 版本号 v1.1.0 -> v1.2.0 ----------
for p in doc.paragraphs:
    set_runs(p, "v1.1.0", "v1.2.0")

doc.save(SRC)
print("SAVED ->", SRC)
