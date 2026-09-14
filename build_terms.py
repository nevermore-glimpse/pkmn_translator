# -*- coding: utf-8 -*-
"""
从 Excel 多语言表提取术语字典。

两种用法：
    · 命令行：python build_terms.py -i 术语表.xlsx -o term_dict.py
    · 被调用：from build_terms import build_terms
              build_terms("术语表.xlsx", "term_dict.py", "英文", "简体中文")
"""
import argparse
import json
import os
import sys

try:
    import openpyxl
except ImportError:
    openpyxl = None


# ============================================================
# 语言别名表
# ============================================================
LANG_ALIASES = {
    "简体中文": ["简体中文", "简中", "中文（简体）", "中文(简体)", "zh-CN", "zh_Hans"],
    "繁体中文": ["繁体中文", "繁中", "中文（繁体）", "中文(繁体)", "zh-TW", "zh_Hant"],
    "英文":   ["英文", "英语", "English", "EN", "Eng"],
    "日文":   ["日文", "日本语", "日语", "Japanese", "JP", "JPN"],
    "西班牙文": ["西班牙文", "西班牙语", "西班牙", "Spanish", "ES", "Español"],
    "德文":   ["德文", "德语", "German", "DE", "Deutsch"],
    "法文":   ["法文", "法语", "French", "FR", "Français"],
    "意大利文": ["意大利文", "意大利语", "Italian", "IT", "Italiano"],
    "韩文":   ["韩文", "韩语", "Korean", "KO", "한국어"],
}


def _log_info(msg, *args):
    try:
        from logger import get_logger
        get_logger("build_terms").info(msg, *args)
    except Exception:
        print(msg % args if args else msg)


def _log_warning(msg, *args):
    try:
        from logger import get_logger
        get_logger("build_terms").warning(msg, *args)
    except Exception:
        print(msg % args if args else msg)


def norm_header(v):
    return "" if v is None else str(v).strip()


def match_lang(header, lang_key):
    h = norm_header(header)
    if not h:
        return False
    aliases = LANG_ALIASES.get(lang_key, [lang_key])
    for a in aliases:
        if h == a:
            return True
    for a in aliases:
        if len(a) >= 3 and a in h:
            return True
    return False


def find_header_row(rows, source_lang, target_lang, max_scan=15):
    for ri, row in enumerate(rows[:max_scan]):
        src_col = tgt_col = None
        for ci, cell in enumerate(row):
            if src_col is None and match_lang(cell, source_lang):
                src_col = ci
            if tgt_col is None and match_lang(cell, target_lang):
                tgt_col = ci
        if src_col is not None and tgt_col is not None:
            return ri, src_col, tgt_col
    return None, None, None


def extract_from_sheet(ws, source_lang, target_lang):
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], None

    ri, sc, tc = find_header_row(rows, source_lang, target_lang)
    if ri is None:
        return [], f"未找到同时含「{source_lang}」和「{target_lang}」的表头行"

    pairs = []
    seen = set()
    sk_empty = sk_same = sk_short = 0

    for row in rows[ri + 1:]:
        if sc >= len(row) or tc >= len(row):
            continue
        s_raw, t_raw = row[sc], row[tc]
        if s_raw is None or t_raw is None:
            sk_empty += 1
            continue
        s, t = str(s_raw).strip(), str(t_raw).strip()
        if not s or not t:
            sk_empty += 1
            continue
        if s == t:
            sk_same += 1
            continue
        if len(s) < 2 or s.isdigit():
            sk_short += 1
            continue
        if s.lower() in seen:
            continue
        seen.add(s.lower())
        pairs.append((s, t))

    stats = {
        "header_row":    ri + 1,
        "src_col":       sc + 1,
        "tgt_col":       tc + 1,
        "total":         len(pairs),
        "skipped_empty": sk_empty,
        "skipped_same":  sk_same,
        "skipped_short": sk_short,
    }
    return pairs, stats


def write_term_dict(all_pairs, out_path, source_lang, target_lang):
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("# -*- coding: utf-8 -*-\n")
        f.write('"""\n')
        f.write("术语表：由 build_terms.py 自动生成\n")
        f.write(f"源语言：{source_lang}    目标语言：{target_lang}\n")
        f.write(f"共 {len(all_pairs)} 条\n")
        f.write('"""\n\n')
        f.write("TERM_DICT = {\n")
        for s, t in all_pairs:
            f.write(f"    {json.dumps(s, ensure_ascii=False)}: "
                    f"{json.dumps(t, ensure_ascii=False)},\n")
        f.write("}\n")


# ============================================================
# 供 commands.py 调用的核心函数
# ============================================================
def list_sheets(excel_path):
    """返回 [(sheet_name, rows, cols), ...]"""
    if openpyxl is None:
        raise RuntimeError("缺少 openpyxl，请执行：pip install openpyxl")
    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    out = [(name, wb[name].max_row, wb[name].max_column)
           for name in wb.sheetnames]
    wb.close()
    return out


def build_terms(excel_path,
                out_path="term_dict.py",
                source_lang="英文",
                target_lang="简体中文",
                sheets=None,
                dedup=True,
                verbose=True):
    """从 Excel 提取术语字典并写入 out_path。返回 dict 结果。"""
    if openpyxl is None:
        return {"ok": False, "total": 0, "out_path": out_path,
                "sheet_stats": [], "errors": ["缺少 openpyxl"]}

    if not os.path.exists(excel_path):
        return {"ok": False, "total": 0, "out_path": out_path,
                "sheet_stats": [], "errors": [f"文件不存在：{excel_path}"]}

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    targets = sheets or wb.sheetnames

    all_pairs = []
    seen_global = set()
    sheet_stats = []
    errors = []

    _log_info("Excel 提取开始：%s  源=%s  译=%s", excel_path, source_lang, target_lang)

    for name in targets:
        if name not in wb.sheetnames:
            errors.append(f"工作表不存在：{name}")
            continue

        ws = wb[name]
        pairs, stats = extract_from_sheet(ws, source_lang, target_lang)

        if stats is None:
            _log_warning("[%s] %s", name, pairs)
            errors.append(f"[{name}] {pairs}")
            continue

        dup = 0
        if dedup:
            new_pairs = []
            for s, t in pairs:
                if s.lower() in seen_global:
                    continue
                seen_global.add(s.lower())
                new_pairs.append((s, t))
            dup = len(pairs) - len(new_pairs)
            pairs = new_pairs

        all_pairs.extend(pairs)
        stats["sheet"] = name
        stats["dup_across_sheets"] = dup
        sheet_stats.append(stats)

        _log_info("[%s] 表头第 %d 行 源列 %d 译列 %d → %d 条%s",
                  name, stats["header_row"], stats["src_col"],
                  stats["tgt_col"], len(pairs),
                  f"  (跨表去重 {dup})" if dup else "")

    wb.close()

    if not all_pairs:
        errors.append("未提取到任何术语对，请检查源/目标语言列名")
        return {"ok": False, "total": 0, "out_path": out_path,
                "sheet_stats": sheet_stats, "errors": errors}

    write_term_dict(all_pairs, out_path, source_lang, target_lang)
    _log_info("Excel 提取完成：%d 条 → %s", len(all_pairs), out_path)

    return {"ok": True, "total": len(all_pairs), "out_path": out_path,
            "sheet_stats": sheet_stats, "errors": errors}


# ============================================================
# CLI 入口
# ============================================================
def main(argv=None):
    ap = argparse.ArgumentParser(
        description="从 Excel 多语言表提取术语字典",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    ap.add_argument("-i", "--input", required=True, help="Excel 文件路径")
    ap.add_argument("-o", "--output", default="term_dict.py", help="输出路径")
    ap.add_argument("--source", default="英文", help="源语言列名")
    ap.add_argument("--target", default="简体中文", help="目标语言列名")
    ap.add_argument("--sheet", action="append", help="只处理指定工作表")
    ap.add_argument("--list-sheets", action="store_true", help="列出所有工作表")
    ap.add_argument("--no-dedup", action="store_true", help="跨表不去重")
    args = ap.parse_args(argv)

    if args.list_sheets:
        for name, r, c in list_sheets(args.input):
            print(f"  {name}  ({r} 行 × {c} 列)")
        return 0

    result = build_terms(
        args.input, args.output, args.source, args.target,
        sheets=args.sheet, dedup=not args.no_dedup,
    )
    for e in result["errors"]:
        print(f"[err] {e}")
    if result["ok"]:
        print(f"\n✔ 共 {result['total']} 条 → {result['out_path']}")
        return 0
    return 1


if __name__ == "__main__":
    sys.exit(main())